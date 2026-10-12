"""Extract raw text from uploaded resumes (PDF, DOCX, TXT)."""
import io
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt"}


class ParseError(ValueError):
    pass


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except (PdfReadError, Exception) as exc:  # noqa: BLE001 - library raises many types
        raise ParseError(f"Could not read this PDF ({exc.__class__.__name__}).") from exc


def _docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ParseError("Could not read this DOCX file.") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ParseError(f"Unsupported file type '{ext or 'unknown'}'. Use PDF, DOCX or TXT.")
    if ext == ".pdf":
        text = _pdf(data)
    elif ext == ".docx":
        text = _docx(data)
    else:
        text = data.decode("utf-8", errors="ignore")
    if len(text.strip()) < 30:
        raise ParseError(
            "No readable text found. Scanned or image-only resumes need OCR before upload."
        )
    return text
